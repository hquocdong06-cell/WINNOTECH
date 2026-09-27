# -*- coding: utf-8 -*-
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os

def generate_8_modules_docx():
    doc = Document()

    COLOR_PRIMARY = RGBColor(31, 78, 120)     # Navy Blue #1F4E78
    COLOR_SECONDARY = RGBColor(46, 117, 182) # Sky Blue #2E75B6
    COLOR_DARK = RGBColor(38, 38, 38)
    COLOR_MUTED = RGBColor(89, 89, 89)

    HEX_PRIMARY = "1F4E78"
    HEX_CALLOUT_BG = "F2F5F8"
    HEX_BORDER = "CCCCCC"

    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    def set_font(run, size=11, bold=False, italic=False, color=COLOR_DARK):
        run.font.name = "Calibri"
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
        run.font.color.rgb = color

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        set_font(r, size=14, bold=True, color=COLOR_PRIMARY)
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(text)
        set_font(r, size=12, bold=True, color=COLOR_SECONDARY)
        return p

    def add_p(text="", bold_prefix=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            set_font(r_pre, size=11, bold=True)
        if text:
            r_txt = p.add_run(text)
            set_font(r_txt, size=11)
        return p

    def add_bullet(text="", bold_prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            set_font(r_pre, size=11, bold=True)
        if text:
            r_txt = p.add_run(text)
            set_font(r_txt, size=11)
        return p

    def add_box(text, title=""):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.8)

        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{HEX_CALLOUT_BG}"/>')
        cell._tc.get_or_add_tcPr().append(shd)

        tcBorders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="none"/>
                <w:left w:val="single" w:sz="24" w:space="0" w:color="{HEX_PRIMARY}"/>
                <w:bottom w:val="none"/>
                <w:right w:val="none"/>
            </w:tcBorders>
        ''')
        cell._tc.get_or_add_tcPr().append(tcBorders)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        if title:
            r_t = p.add_run(f"{title}\n")
            set_font(r_t, size=11, bold=True, color=COLOR_PRIMARY)
        r = p.add_run(text)
        set_font(r, size=10.5, italic=True)
        doc.add_paragraph()

    # Title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("SỔ TAY BẢO VỆ ĐỒ ÁN: TOÀN BỘ 8 PHÂN HỆ WINNOTECH\n(GIẢI THÍCH NGHIỆP VỤ - DB LƯU Ở ĐÂU - CODE Ở ĐÂU)")
    set_font(r, size=17, bold=True, color=COLOR_PRIMARY)

    add_box(
        "Tài liệu này bám sát 100% danh mục 8 phân hệ chức năng thực tế của dự án WINNOTech.\n"
        "Mỗi chức năng được làm rõ 3 câu hỏi sống còn khi bảo vệ:\n"
        "1. Nghiệp vụ hoạt động ra sao (luồng đi dữ liệu)?\n"
        "2. Database lưu ở bảng nào (tên Collection & các trường chính)?\n"
        "3. Mã nguồn nằm ở file nào (Frontend UI & Backend API)?",
        "🎯 MỤC TIÊU: HIỂU TRỌN VẸN ĐỒ ÁN - ĐỖ MÔN ĐIỂM CAO"
    )

    # =========================================================================
    # PHÂN HỆ 1
    # =========================================================================
    add_h1("PHÂN HỆ 1: NGƯỜI DÙNG & XÁC THỰC (AUTHENTICATION & PROFILE)")

    add_h2("1.1 Đăng ký & Đăng nhập (JWT + Bcrypt)")
    add_bullet("Người dùng điền email, mật khẩu. Server dùng bcrypt băm mật khẩu (10 vòng salt). Khi đăng nhập, server kiểm tra password bằng bcrypt.compare, nếu đúng thì sinh mã JWT token (ký thuật toán RS256 hoặc HS256) gửi về cho trình duyệt lưu vào localStorage.", "Nghiệp vụ: ")
    add_bullet("Bảng 'User' (models/User.js). Các trường: name, email, password, role ('user'|'admin'), status ('active'|'blocked').", "DB lưu ở đâu: ")
    add_bullet("Frontend: frontend/src/pages/Auth.jsx | Backend: server.js (POST /api/register, POST /api/login).", "Code ở đâu: ")

    add_h2("1.2 Đăng nhập bằng Google OAuth 2.0")
    add_bullet("Người dùng bấm 'Đăng nhập với Google' ➔ Google cấp credential id_token ➔ Frontend gửi token này lên server ➔ Server dùng thư viện google-auth-library để verify ➔ Nếu chưa có tài khoản thì tự động tạo User mới trong DB với googleId.", "Nghiệp vụ: ")
    add_bullet("Bảng 'User' (trường googleId, email, name, avatar).", "DB lưu ở đâu: ")
    add_bullet("Frontend: Auth.jsx | Backend: server.js (POST /api/google-login).", "Code ở đâu: ")

    add_h2("1.3 Quên mật khẩu & Gửi OTP qua Email (Nodemailer)")
    add_bullet("Người dùng nhập email ➔ Server sinh ngẫu nhiên mã OTP 6 số (crypto.randomInt) và lưu vào User kèm thời hạn hết hạn 10 phút ➔ Gửi OTP qua Gmail bằng Nodemailer ➔ Người dùng nhập đúng OTP thì mới được đổi mật khẩu mới.", "Nghiệp vụ: ")
    add_bullet("Bảng 'User' (trường resetPasswordOTP, resetPasswordExpires).", "DB lưu ở đâu: ")
    add_bullet("Frontend: Auth.jsx | Backend: server.js (POST /api/forgot-password, /api/verify-otp, /api/reset-password).", "Code ở đâu: ")

    add_h2("1.4 Quản lý Hồ sơ & Sổ địa chỉ giao hàng")
    add_bullet("Người dùng cập nhật họ tên, SĐT, ảnh đại diện avatar (upload qua Multer). Có thể thêm nhiều địa chỉ giao hàng và tích chọn 1 địa chỉ làm mặc định.", "Nghiệp vụ: ")
    add_bullet("Bảng 'User' (thông tin cá nhân) và Bảng 'DeliveryAddress' (models/DeliveryAddress.js: u_id, name, phone, address, is_default).", "DB lưu ở đâu: ")
    add_bullet("Frontend: frontend/src/pages/Profile.jsx | Backend: server.js (/api/profile, /api/delivery-addresses).", "Code ở đâu: ")

    add_h2("1.5 Quản lý Đơn hàng cá nhân")
    add_bullet("Xem lịch sử mua hàng, trạng thái đơn (Chờ xử lý ➔ Đang giao ➔ Hoàn thành / Hủy), xem từng linh kiện đã mua và tổng tiền.", "Nghiệp vụ: ")
    add_bullet("Bảng 'Order' (user_id, code, total_amount, status) và Bảng 'OrderItem' (order_id, variants_id, Quantity, price).", "DB lưu ở đâu: ")
    add_bullet("Frontend: Profile.jsx (tab Đơn hàng) | Backend: server.js & routers/order.js (/api/my-orders).", "Code ở đâu: ")

    # =========================================================================
    # PHÂN HỆ 2
    # =========================================================================
    add_h1("PHÂN HỆ 2: MUA SẮM & BỘ LỌC SẢN PHẨM (CATALOG & SMART FILTERS)")

    add_h2("2.1 Trang chủ (Home Page) & Flash Sale")
    add_bullet("Hiển thị slider banner quảng cáo, đồng hồ đếm ngược Flash Sale thời gian thực, danh sách linh kiện bán chạy (sắp xếp theo sold_count giảm dần).", "Nghiệp vụ: ")
    add_bullet("Bảng 'Banner' (models/BannerPaymentImage.js), Bảng 'Product', Bảng 'FlashSaleSetting'.", "DB lưu ở đâu: ")
    add_bullet("Frontend: frontend/src/pages/Home.jsx | Backend: server.js (/api/banners, /api/flash-sale, /api/products).", "Code ở đâu: ")

    add_h2("2.2 Bộ lọc thông minh theo danh mục (Smart Filter)")
    add_bullet("Khách hàng lọc theo Hãng sản xuất (ASUS, MSI, Gigabyte...), theo mức giá (min/max), và theo thông số phần cứng riêng biệt (Socket LGA1700/AM5 cho CPU; Chipset B760/Z790 cho Mainboard; Bus RAM, Chuẩn giao tiếp NVMe, Công suất nguồn PSU).", "Nghiệp vụ: ")
    add_bullet("Bảng 'Category', Bảng 'Brand', Bảng 'Specification' (models/Specification.js: p_id, id_attribute_value) liên kết với Bảng 'AttributeValue'.", "DB lưu ở đâu: ")
    add_bullet("Frontend: frontend/src/pages/CategoryPage.jsx, CPU.jsx, GPU.jsx, RAM.jsx, Mainboard.jsx, PriceRangeFilter.jsx | Backend: server.js.", "Code ở đâu: ")

    add_h2("2.3 Tìm kiếm thông minh (Smart Search với Fuse.js)")
    add_bullet("Khách gõ từ khóa vào ô tìm kiếm ➔ Thư viện Fuse.js thực hiện tìm kiếm mờ (Fuzzy Search) theo tên sản phẩm, thương hiệu ngay khi gõ và hiển thị dropdown kết quả tức thì.", "Nghiệp vụ: ")
    add_bullet("Dữ liệu lấy từ Bảng 'Product'.", "DB lưu ở đâu: ")
    add_bullet("Frontend: frontend/src/layouts/Header.jsx (hoặc Navbar) | Backend: server.js (/api/search).", "Code ở đâu: ")

    add_h2("2.4 Trang Chi tiết sản phẩm (Product Detail)")
    add_bullet("Hiển thị bộ ảnh gallery có thể chọn xem, bảng thông số kỹ thuật (Specification), và quan trọng nhất: cho phép khách CHỌN BIẾN THỂ (ví dụ chọn màu Đen hay Trắng, bản 8GB hay 16GB). Khi chọn biến thể, giá bán (price, sale_price) và số lượng tồn kho (stock_quantity) tự động cập nhật.", "Nghiệp vụ: ")
    add_bullet("Bảng 'Product' (thông tin cha) liên kết 1-n với Bảng 'ProductVariant' (các phiên bản con) và Bảng 'variants_attributes'.", "DB lưu ở đâu: ")
    add_bullet("Frontend: frontend/src/pages/ProductDetail.jsx | Backend: server.js (/api/product/:slug).", "Code ở đâu: ")

    # =========================================================================
    # PHÂN HỆ 3
    # =========================================================================
    add_h1("PHÂN HỆ 3: XÂY DỰNG CẤU HÌNH PC (PC BUILDER & TƯƠNG THÍCH)")

    add_h2("3.1 Xây dựng cấu hình 8 nhóm linh kiện thiết yếu")
    add_bullet("Khách hàng tự chọn từng linh kiện theo đúng quy trình: 1. CPU ➔ 2. Mainboard ➔ 3. RAM ➔ 4. Ổ cứng (SSD/HDD) ➔ 5. Card đồ họa (VGA) ➔ 6. Nguồn (PSU) ➔ 7. Tản nhiệt ➔ 8. Vỏ Case.", "Nghiệp vụ: ")
    add_bullet("Bảng 'BuildPC' & 'BuildItem' (models/BuildPc.js) liên kết với 'Product' và 'ProductVariant'.", "DB lưu ở đâu: ")
    add_bullet("Frontend: frontend/src/pages/BuildPC.jsx | Backend: server.js (/api/build-pc/...).", "Code ở đâu: ")

    add_h2("3.2 Thuật toán kiểm tra độ tương thích (Compatibility Check)")
    add_bullet("Hệ thống kiểm tra thông số kỹ thuật (Specification): Ví dụ CPU dùng socket 'LGA1700' thì chỉ cho phép chọn Mainboard có socket 'LGA1700'; Mainboard chuẩn RAM 'DDR5' thì chỉ ghép được với thanh RAM 'DDR5'. Đồng thời tự động cộng tổng công suất (Watt) của CPU + VGA để cảnh báo chọn Nguồn (PSU) đủ công suất.", "Nghiệp vụ: ")
    add_bullet("Dữ liệu đối chiếu từ Bảng 'Specification' và 'AttributeValue'.", "DB lưu ở đâu: ")
    add_bullet("Frontend: BuildPC.jsx | Backend: scripts kiểm tra tương thích trong server.js.", "Code ở đâu: ")

    add_h2("3.3 Xuất file báo giá & Thêm cả dàn PC vào giỏ hàng")
    add_bullet("Khách có thể bấm 'Xuất PDF/Excel' để in bảng báo giá cấu hình máy tính mang đi tư vấn; hoặc bấm 'Thêm vào giỏ hàng' để đẩy cùng lúc cả 8 linh kiện vào giỏ để thanh toán ngay.", "Nghiệp vụ: ")
    add_bullet("Lưu tạm vào Redux và chuyển thành các bản ghi trong Bảng 'CartItem'.", "DB lưu ở đâu: ")
    add_bullet("Frontend: BuildPC.jsx | Backend: server.js.", "Code ở đâu: ")

    # =========================================================================
    # PHÂN HỆ 4
    # =========================================================================
    add_h1("PHÂN HỆ 4: GIỎ HÀNG & THANH TOÁN (CART, VOUCHER & CHECKOUT)")

    add_h2("4.1 Giỏ hàng thông minh (Dual-State Cart)")
    add_bullet("Khi chưa đăng nhập: Giỏ hàng lưu tạm ở LocalStorage trình duyệt bằng Redux Toolkit. Khi đăng nhập: Hệ thống tự động đồng bộ gộp các món từ LocalStorage vào Database MongoDB để lưu vĩnh viễn.", "Nghiệp vụ: ")
    add_bullet("Bảng 'CartItem' (u_id, variants_id, Quantity).", "DB lưu ở đâu: ")
    add_bullet("Frontend: frontend/src/redux/cartSlice.js, Cart.jsx | Backend: server.js (/api/cart).", "Code ở đâu: ")

    add_h2("4.2 Hệ thống Mã giảm giá Voucher đa cấp")
    add_bullet("Khách nhập mã khuyến mãi. Mã có chữ 'FRS' ➔ Miễn phí vận chuyển 100% + giảm tiền hàng. Mã có chữ 'SHIP' ➔ Chỉ giảm tiền ship, không trừ tiền sản phẩm. Mã thường ➔ Trừ tiền hàng theo % hoặc số tiền cố định. Có kiểm tra điều kiện đơn hàng tối thiểu (min_order_value) và hạn dùng.", "Nghiệp vụ: ")
    add_bullet("Bảng 'Voucher' và 'UserVoucher' (models/Voucher.js).", "DB lưu ở đâu: ")
    add_bullet("Frontend: Checkout.jsx | Backend: server.js (hàm calculateVoucherDiscount dòng 93).", "Code ở đâu: ")

    add_h2("4.3 Thanh toán COD & Cổng trực tuyến VNPay Sandbox")
    add_bullet("COD: Nhận hàng trả tiền mặt. VNPay: Server tạo link thanh toán và mã QR với chữ ký băm bảo mật HMAC-SHA512. Khi khách quét mã qua app ngân hàng thành công, ngân hàng gọi webhook IPN và điều hướng về /payment-result ➔ Hệ thống cập nhật đơn thành 'paid' và trừ tồn kho trong DB.", "Nghiệp vụ: ")
    add_bullet("Bảng 'Order' (payment_status: 'unpaid'|'paid') và Bảng 'OrderItem'.", "DB lưu ở đâu: ")
    add_bullet("Frontend: Checkout.jsx, PaymentResult.jsx | Backend: routers/order.js (/order/create_payment_url, /order/vnpay_return).", "Code ở đâu: ")

    # =========================================================================
    # PHÂN HỆ 5
    # =========================================================================
    add_h1("PHÂN HỆ 5: ĐÁNH GIÁ, YÊU THÍCH & SO SÁNH (SOCIAL & ENGAGEMENT)")

    add_h2("5.1 Đánh giá & Bình luận (Review & Rating)")
    add_bullet("Khách chấm sao từ 1 đến 5 sao, viết nhận xét và đính kèm ảnh chụp linh kiện thực tế. Chỉ những khách hàng đã mua sản phẩm đó (có liên kết id_orderitems) mới được hiển thị huy hiệu 'Đã mua hàng tại WINNOTech'.", "Nghiệp vụ: ")
    add_bullet("Bảng 'Review' (models/FavoriteCompareReview.js: user_id, product_id, id_orderitems, rating, comment, images, status: 'approved').", "DB lưu ở đâu: ")
    add_bullet("Frontend: ProductDetail.jsx | Backend: server.js (/api/reviews).", "Code ở đâu: ")

    add_h2("5.2 Sản phẩm yêu thích (Wishlist)")
    add_bullet("Khách bấm vào icon trái tim ở sản phẩm để lưu vào danh sách yêu thích, giúp xem lại nhanh các linh kiện muốn mua mà chưa có tiền ngay.", "Nghiệp vụ: ")
    add_bullet("Bảng 'Favorite' (models/FavoriteCompareReview.js: user_id, product_id).", "DB lưu ở đâu: ")
    add_bullet("Frontend: ProductCard.jsx, Profile.jsx | Backend: server.js (/api/favorites).", "Code ở đâu: ")

    add_h2("5.3 So sánh linh kiện (Compare)")
    add_bullet("Khách chọn 2 hoặc nhiều linh kiện cùng danh mục (ví dụ 2 card màn hình RTX 4060 và RTX 3060) ➔ Hệ thống hiển thị bảng đối chiếu từng dòng thông số kỹ thuật (VRAM, Bus, Cổng kết nối, Kích thước) song song để khách dễ chọn mua.", "Nghiệp vụ: ")
    add_bullet("Bảng 'Compare' và Bảng 'Specification'.", "DB lưu ở đâu: ")
    add_bullet("Frontend: frontend/src/pages/Compare.jsx | Backend: server.js (/api/compare).", "Code ở đâu: ")

    # =========================================================================
    # PHÂN HỆ 6
    # =========================================================================
    add_h1("PHÂN HỆ 6: TIN TỨC & THỦ THUẬT CÔNG NGHỆ (BLOG & NEWS)")

    add_h2("6.1 Danh mục bài viết & Bài viết nổi bật")
    add_bullet("Hiển thị tin công nghệ, bài review đánh giá phần cứng, cẩm nang hướng dẫn build máy tính. Có phân mục rõ ràng, phân trang và hiển thị bài đọc nhiều nhất.", "Nghiệp vụ: ")
    add_bullet("Bảng 'Post' & Bảng 'PostCategory' (models/Post.js: title, slug, content, thumbnail, views, category_id).", "DB lưu ở đâu: ")
    add_bullet("Frontend: frontend/src/pages/Blog.jsx | Backend: server.js (/api/posts).", "Code ở đâu: ")

    add_h2("6.2 Chi tiết bài viết & Gợi ý sản phẩm liên quan")
    add_bullet("Nội dung bài viết định dạng phong phú (Rich Text/HTML). Bên dưới bài viết tự động gợi ý các linh kiện máy tính được nhắc đến trong bài để kích thích khách bấm mua.", "Nghiệp vụ: ")
    add_bullet("Bảng 'Post' liên kết với Bảng 'Product'.", "DB lưu ở đâu: ")
    add_bullet("Frontend: frontend/src/pages/BlogPostDetail.jsx | Backend: server.js (/api/post/:slug).", "Code ở đâu: ")

    # =========================================================================
    # PHÂN HỆ 7
    # =========================================================================
    add_h1("PHÂN HỆ 7: QUẢN TRỊ HỆ THỐNG (ADMIN MANAGEMENT PANEL)")

    add_h2("7.1 Dashboard Thống kê doanh thu")
    add_bullet("Hiển thị tổng doanh thu, số đơn thành công, khách hàng mới. Dùng Aggregation Pipeline gom nhóm doanh thu theo ngày/tháng để vẽ biểu đồ trực quan (React Google Charts).", "Nghiệp vụ: ")
    add_bullet("Query tổng hợp từ Bảng 'Order' (status='completed') và Bảng 'User'.", "DB lưu ở đâu: ")
    add_bullet("Frontend: frontend/src/admin/pages/Dashboard.jsx | Backend: server.js (/api/admin/dashboard-stats, /api/admin/revenue-chart).", "Code ở đâu: ")

    add_h2("7.2 Quản lý Sản phẩm, Biến thể & Tồn kho")
    add_bullet("Admin thêm sản phẩm cha, thêm các biến thể con (SKU, giá bán, giá khuyến mãi, số lượng tồn kho), tải ảnh đại diện lên server, cập nhật số lượng nhập kho.", "Nghiệp vụ: ")
    add_bullet("Bảng 'Product' và Bảng 'ProductVariant'.", "DB lưu ở đâu: ")
    add_bullet("Frontend: frontend/src/admin/pages/Products.jsx | Backend: server.js (/api/admin/products).", "Code ở đâu: ")

    add_h2("7.3 Quản lý Đơn hàng & In hóa đơn, Xuất Excel")
    add_bullet("Admin xem danh sách đơn hàng, lọc theo trạng thái (Chờ xác nhận, Đang chuẩn bị, Đang giao, Đã giao, Đã hủy). Admin đổi trạng thái thì hệ thống bắn Socket.IO báo về máy khách hàng. Có nút xuất báo cáo Excel (ExcelJS) và In hóa đơn vận chuyển PDF (PDFKit).", "Nghiệp vụ: ")
    add_bullet("Bảng 'Order' (cập nhật status, statusHistory) và Bảng 'OrderItem'.", "DB lưu ở đâu: ")
    add_bullet("Frontend: frontend/src/admin/pages/Orders.jsx | Backend: server.js & routers/order.js.", "Code ở đâu: ")

    add_h2("7.4 Quản lý Khuyến mãi, Bài viết, Đánh giá & Người dùng")
    add_bullet("Tạo mã voucher, thiết lập Flash Sale; duyệt hoặc ẩn các bình luận đánh giá; viết bài đăng tin tức; xem danh sách khách hàng và khóa/mở khóa tài khoản vi phạm.", "Nghiệp vụ: ")
    add_bullet("Bảng 'Voucher', 'Post', 'Review', 'User'.", "DB lưu ở đâu: ")
    add_bullet("Frontend: Promotions.jsx, Posts.jsx, Reviews.jsx, Customers.jsx | Backend: server.js.", "Code ở đâu: ")

    # =========================================================================
    # PHÂN HỆ 8 & KẾT LUẬN
    # =========================================================================
    add_h1("PHÂN HỆ 8: NỀN TẢNG CÔNG NGHỆ (TECH STACK TỔNG HỢP)")
    add_bullet("React.js 18, Vite, Redux Toolkit, React Router DOM v7, TailwindCSS, Socket.io-client, Lucide React, Fuse.js, React Google Charts.", "Frontend: ")
    add_bullet("Node.js, Express.js 5 RESTful API, JWT RS256, Bcrypt, Nodemailer, VNPay SDK, PDFKit, ExcelJS, Socket.IO, Google Generative AI (Gemini).", "Backend: ")
    add_bullet("MongoDB Local cổng 27017 (hoặc MongoDB Atlas Cloud) + Mongoose ORM.", "Cơ sở dữ liệu: ")

    add_box(
        "KHI ĐI THI BẢO VỆ:\n"
        "• Thầy cô hỏi CSDL: Mở MongoDB Compass chỉ vào DB 'WINNOTech'.\n"
        "• Thầy cô bảo mở Code: Mở VS Code, cần chức năng nào thì tra đúng tên file ở trên là mở trúng 100%!\n"
        "• Thầy cô hỏi luồng: Nêu 3 bước: Khách thao tác trên React ➔ Gọi API Express ➔ Truy vấn MongoDB và trả kết quả.",
        "💡 BẢO BỐI PHÒNG THI"
    )

    out_name = "SO_TAY_BAO_VE_TOAN_BO_8_PHAN_HE_WINNOTECH.docx"
    out_path = os.path.join(os.path.dirname(__file__), out_name)
    doc.save(out_path)
    print("Done:", out_path)

if __name__ == "__main__":
    generate_8_modules_docx()
