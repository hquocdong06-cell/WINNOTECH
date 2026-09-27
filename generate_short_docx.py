# -*- coding: utf-8 -*-
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os

def create_short_doc():
    doc = Document()

    COLOR_PRIMARY = RGBColor(31, 78, 120)     # Navy Blue
    COLOR_DARK = RGBColor(40, 40, 40)
    HEX_PRIMARY = "1F4E78"
    HEX_CALLOUT_BG = "F2F5F8"
    HEX_BORDER = "CCCCCC"

    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    def set_font(run, size=11, bold=False, italic=False, color=COLOR_DARK):
        run.font.name = "Calibri"
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
        run.font.color.rgb = color

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        set_font(r, size=14, bold=True, color=COLOR_PRIMARY)
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(text)
        set_font(r, size=12, bold=True, color=COLOR_PRIMARY)
        return p

    def add_p(text="", bold_prefix=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            set_font(r_pre, size=11, bold=True)
        if text:
            r_txt = p.add_run(text)
            set_font(r_txt, size=11)
        return p

    def add_bullet(text="", bold_prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            set_font(r_pre, size=11, bold=True)
        if text:
            r_txt = p.add_run(text)
            set_font(r_txt, size=11)
        return p

    def add_box(text, title=""):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.8)

        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{HEX_CALLOUT_BG}"/>')
        cell._tc.get_or_add_tcPr().append(shd)

        tcBorders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="none"/>
                <w:left w:val="single" w:sz="24" w:space="0" w:color="{HEX_PRIMARY}"/>
                <w:bottom w:val="none"/>
                <w:right w:val="none"/>
            </w:tcBorders>
        ''')
        cell._tc.get_or_add_tcPr().append(tcBorders)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        if title:
            r_t = p.add_run(f"{title}\n")
            set_font(r_t, size=11, bold=True, color=COLOR_PRIMARY)
        r = p.add_run(text)
        set_font(r, size=10.5, italic=True)
        doc.add_paragraph()

    # --- TITLE ---
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("TÀI LIỆU ÔN NHANH DỰ ÁN WINNOTECH\n(HIỂU NGHIỆP VỤ & DATABASE ĐỂ BẢO VỆ ĐỖ MÔN)")
    set_font(r, size=18, bold=True, color=COLOR_PRIMARY)

    add_box(
        "Bản tóm tắt ngắn gọn nhất: Không lý thuyết dài dòng, chỉ tập trung vào 2 điều cốt lõi:\n"
        "1. Cơ sở dữ liệu lưu ở đâu? Bảng nào chứa cái gì?\n"
        "2. Nghiệp vụ từng chức năng chạy như thế nào (kèm file code để chỉ cho thầy cô)?",
        "📌 DÀNH CHO BẠN TRƯỚC GIỜ BẢO VỆ"
    )

    # --- PHẦN 1 ---
    add_h1("PHẦN 1: DATABASE LƯU Ở ĐÂU? (CỰC KỲ DỄ HIỂU)")
    add_p("Hệ thống dùng cơ sở dữ liệu MongoDB Local (chạy trên máy của bạn).")
    add_bullet("mongodb://127.0.0.1:27017/WINNOTech (xem file .env dòng 11)", "Đường link kết nối: ")
    add_bullet("WINNOTech", "Tên Database: ")
    add_bullet("Mở phần mềm MongoDB Compass, bấm Connect là thấy ngay database WINNOTech.", "Cách mở xem: ")

    add_h2("Dữ liệu lưu vào 6 Bảng (Collection) chính sau:")
    add_bullet("Lưu tài khoản người dùng, email, mật khẩu đã mã hóa (bcrypt), role ('admin' hoặc 'user'), số dư tiền ví.", "1. Bảng 'User': ")
    add_bullet("Lưu thông tin sản phẩm chung: tên, ảnh, danh mục (CPU, VGA, RAM...), hãng (ASUS, MSI...), mô tả.", "2. Bảng 'Product': ")
    add_bullet("Lưu từng phiên bản của sản phẩm (ví dụ bản 8GB, 16GB, Đen, Trắng). Mỗi biến thể có: Giá bán (price), Giá giảm (sale_price), và Số lượng còn trong kho (stock_quantity).", "3. Bảng 'ProductVariant': ")
    add_bullet("Lưu giỏ hàng khi người dùng bấm 'Thêm vào giỏ'. Gồm: ID người dùng, ID biến thể sản phẩm, số lượng mua.", "4. Bảng 'CartItem': ")
    add_bullet("Lưu thông tin đơn hàng: Mã đơn (code), tên khách, SĐT, địa chỉ, tổng tiền, trạng thái đơn ('pending', 'shipping', 'completed', 'cancelled'), vết lịch sử đơn (statusHistory).", "5. Bảng 'Order': ")
    add_bullet("Lưu chi tiết từng món đồ trong đơn: Mua biến thể nào, số lượng bao nhiêu, giá tiền lúc mua là bao nhiêu (Snapshot giá).", "6. Bảng 'OrderItem': ")

    add_box(
        "Thầy cô hỏi: 'Nếu sau này xóa sản phẩm hoặc đổi giá thì đơn hàng cũ có bị đổi giá theo không?'\n"
        "Trả lời ngay: 'Dạ KHÔNG, vì khi đặt hàng, giá mua đã được lưu cố định vào bảng OrderItem rồi (kỹ thuật Snapshot), nên dù sản phẩm gốc có đổi giá thì lịch sử đơn hàng cũ vẫn giữ nguyên tiền ạ.'",
        "⭐ CÂU HỎI THẦY CÔ HAY BẮT BẺ VỀ DATABASE"
    )

    # --- PHẦN 2 ---
    add_h1("PHẦN 2: NGHIỆP VỤ CÁC CHỨC NĂNG (LUỒNG CHẠY & FILE CODE)")

    add_h2("1. Nghiệp vụ Mua hàng & Giỏ hàng (Cart & Checkout)")
    add_bullet("Khách bấm 'Thêm vào giỏ' ➔ Lưu vào bảng 'CartItem' trong DB. Nếu chưa đăng nhập thì Redux lưu tạm trên giao diện.", "• Cách chạy: ")
    add_bullet("Khi bấm Đặt hàng (Checkout) ➔ Tạo 1 đơn mới trong bảng 'Order', tạo các món đồ trong bảng 'OrderItem', đồng thời tự động TRỪ tồn kho trong 'ProductVariant' (stock_quantity giảm đi).", "• Lúc đặt hàng: ")
    add_bullet("Frontend: frontend/src/pages/Cart.jsx & Checkout.jsx | Backend: server.js", "• Code ở đâu: ")

    add_h2("2. Nghiệp vụ Giảm giá / Voucher (FRS & SHIP)")
    add_bullet("Voucher có chữ FRS (như FRS50K): Miễn phí ship 100% + giảm thêm tiền sản phẩm.", "• Voucher FRS: ")
    add_bullet("Voucher có chữ SHIP (như SHIP20): Chỉ trừ tiền vận chuyển, TUYỆT ĐỐI không trừ vào tiền sản phẩm.", "• Voucher SHIP: ")
    add_bullet("Voucher thường: Giảm tiền sản phẩm, tiền ship vẫn tính bình thường.", "• Voucher thường: ")
    add_bullet("File server.js, tìm hàm calculateVoucherDiscount (dòng 93).", "• Code ở đâu: ")

    add_h2("3. Nghiệp vụ Thanh toán VNPay (Quét mã QR / Thẻ)")
    add_bullet("Khách chọn thanh toán VNPay ➔ Server tạo 1 đường link kèm mã QR có chữ ký bảo mật băm HMAC-SHA512 để không ai sửa được số tiền.", "• Cách chạy: ")
    add_bullet("Khách quét mã xong ➔ Ngân hàng trả kết quả về trang /payment-result ➔ Nếu mã trả về là '00' thì hệ thống đổi trạng thái đơn thành 'Đã thanh toán' (paid).", "• Hoàn tất: ")
    add_bullet("File routers/order.js & frontend/src/pages/PaymentResult.jsx", "• Code ở đâu: ")

    add_h2("4. Nghiệp vụ Lịch sử Đơn hàng & Hủy đơn hoàn kho")
    add_bullet("Đơn hàng đi qua các bước: Chờ xác nhận (pending) ➔ Chuẩn bị hàng (preparing) ➔ Đang giao (shipping) ➔ Đã giao (delivered) ➔ Hoàn thành (completed).", "• Vòng đời đơn: ")
    add_bullet("Khách hoặc Admin chỉ được hủy khi đơn ở bước 'Chờ xác nhận' hoặc 'Chuẩn bị hàng'. Khi ấn Hủy đơn, hệ thống TỰ ĐỘNG CỘNG LẠI số lượng tồn kho ($inc stock_quantity) trong bảng ProductVariant.", "• Hủy đơn hoàn kho: ")
    add_bullet("File routers/order.js & server.js", "• Code ở đâu: ")

    add_h2("5. Nghiệp vụ Chatbot AI tư vấn cấu hình máy tính")
    add_bullet("Khách gõ: 'Tư vấn cho tôi dàn máy 15 triệu chơi game'.", "• Khách hỏi: ")
    add_bullet("Hệ thống tự bóc tách số tiền '15 triệu', sau đó TỰ VÀO MONGODB TÌM các linh kiện thật có giá dưới 15 triệu, rồi nạp vào cho Gemini AI trả lời. Vì vậy AI không bịa linh kiện vớ vẩn mà đề xuất đúng đồ cửa hàng đang có bán!", "• Cách AI chạy: ")
    add_bullet("File routers/AI_chatbot.js & frontend/src/components/AIChatbot.jsx", "• Code ở đâu: ")

    add_h2("6. Nghiệp vụ Admin Dashboard & Báo cáo")
    add_bullet("Thống kê doanh thu theo ngày/tháng, vẽ biểu đồ (React Google Charts).", "• Thống kê: ")
    add_bullet("Có nút xuất file Excel doanh thu (dùng thư viện exceljs) và in hóa đơn giao hàng PDF (dùng thư viện pdfkit).", "• Xuất báo cáo: ")
    add_bullet("File frontend/src/admin/pages/Dashboard.jsx & server.js", "• Code ở đâu: ")

    # --- PHẦN 3 ---
    add_h1("PHẦN 3: BÍ QUYẾT BẢO VỆ CHỐNG TRƯỢT (NÓI GÌ KHI ĐỨNG TRƯỚC THẦY CÔ?)")
    add_bullet("1. Thưa Thầy/Cô, dự án của em là web bán linh kiện máy tính WINNOTech, điểm mạnh là có quản lý đa biến thể sản phẩm, thanh toán VNPay và Chatbot AI tư vấn PC theo ngân sách.", "Câu 1: ")
    add_bullet("2. Cơ sở dữ liệu em dùng MongoDB, lưu ở máy cục bộ cổng 27017, gồm các bảng chính là User, Product, ProductVariant, Order, OrderItem.", "Câu 2: ")
    add_bullet("3. Khi khách hủy đơn, hệ thống tự động hoàn lại tồn kho cho cửa hàng.", "Câu 3: ")
    add_bullet("4. Nếu thầy cô bảo mở code: Bấm Alt+Tab qua VS Code, mở file server.js (xử lý logic chính) hoặc routers/order.js (đơn hàng & VNPay).", "Câu 4: ")

    out_name = "HUONG_DAN_NGAN_GON_NGHIEP_VU_VA_DB_WINNOTECH.docx"
    out_path = os.path.join(os.path.dirname(__file__), out_name)
    doc.save(out_path)
    print("Done:", out_path)

if __name__ == "__main__":
    create_short_doc()
